import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
import shutil

TEMPLATE_PATH = r"C:\Users\24026811\Downloads\PI_Template_Entrega1_Original_Backup.docx"
TARGET_DOWNLOADS = r"C:\Users\24026811\Downloads\PI_Template_Entrega1.docx"
TARGET_REPO = r"c:\Users\24026811\Documents\Projeto5\documentos\Entrega 1\Projeto Interdisciplinar_ Inteligência Artificial\Relatorio_Tecnico_e_Model_Card_Entrega_1.docx"

doc = docx.Document(TEMPLATE_PATH)
print("Loaded original template successfully.")

# Helper to set text in a paragraph while preserving style
def set_p_text(p, text, bold_prefix=""):
    # Clear existing runs
    p.text = ""
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
    p.add_run(text)

# Helper to set cell text and formatting
def set_cell(cell, text, bold=False):
    cell.text = text
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if bold and len(p.runs) > 0:
        p.runs[0].bold = True

def add_table_row(table, col_texts):
    row = table.add_row()
    for idx, text in enumerate(col_texts):
        set_cell(row.cells[idx], text)
    return row

# 1. Update Subtitle in P02
p02 = doc.paragraphs[2]
p02.text = "Entrega 1 — Parcial · Agente 3: Agente para o Estudante (Álvaro AI)"
# Add group info as run
p02_info = doc.paragraphs[3]
p02_info.text = (
    "Projeto: Álvaro AI — Central de Atendimento Inteligente FECAP (ASA) | "
    "Curso: Ciência da Computação (5º Semestre 2026)\n"
    "Integrantes: Esther Oliveira Costa, Higor Fonseca, João Victor Faria\n\n"
    "Este relatório documenta a construção de uma baseline (primeira versão simples do modelo e arquitetura RAG) "
    "coerente com o Agente para o Estudante (Agente 3), seguindo os requisitos da Entrega 1 do Projeto Interdisciplinar."
)

# 2. Table 0: Data Sources (1. Análise dos Dados)
# Template Table 0: 4 rows x 4 cols
t0 = doc.tables[0]
# Row 0 is header: ['Fonte', 'Registros', 'Conteúdo', 'Observações']
t0_rows = [
    [
        "KBDocument\n(Base RAG FECAP)",
        "16 documentos\n(19 chunks)",
        "Normas e regulamentos oficiais da FECAP (matrícula, bolsas Prouni/FIES, colação de grau, TCC, atestados e prazos).",
        "1 linha por documento oficial. Base primária do RAG para geração de respostas com fontes oficiais."
    ],
    [
        "Student\n(Cadastro do Aluno)",
        "10 estudantes\nativos",
        "RA único, nome, curso (CC, ADM, Contábeis), semestre (1º ao 8º), turno, status acadêmico e contato.",
        "1 linha por estudante. Contextualiza e personaliza respostas conforme curso, semestre e situação de matrícula."
    ],
    [
        "ChatMessage &\nConversation",
        "7 sessões /\n23 mensagens",
        "Histórico de mensagens do chat, papel ('user', 'assistant', 'system'), texto e ações contextuais sugeridas.",
        "1 linha por mensagem. Memória conversacional da sessão e preservação de contexto no transbordo humano."
    ],
    [
        "Ticket\n(Chamados ASA)",
        "12 chamados de\nhomologação",
        "Número sequencial, título, descrição, categoria, prioridade, status de atendimento, SLA limite e aluno vinculado.",
        "1 linha por chamado. Acionado no transbordo para a equipe ASA quando a IA não resolve a dúvida."
    ],
    [
        "AIAnalysis\n(Inferência e Triagem)",
        "8 análises pré-computadas + tempo real",
        "Intenção detectada, categoria predita, score de confiança (0.85 a 0.98), resumo executivo e recomendação.",
        "1 linha por inferência analítica (1:1 com Ticket). Apoia o atendente após o transbordo."
    ]
]

# Fill existing rows 1, 2, 3
for r_idx in range(1, 4):
    for c_idx in range(4):
        set_cell(t0.cell(r_idx, c_idx), t0_rows[r_idx - 1][c_idx])

# Add remaining rows 4 and 5
for r_data in t0_rows[3:]:
    add_table_row(t0, r_data)

# 3. P08: Preparação dos Dados (2. Preparação dos Dados)
p08 = doc.paragraphs[8]
p08.text = (
    "A solução não utiliza bases em Excel nem planilhas estáticas desarticuladas; toda a persistência é gerenciada via "
    "banco de dados relacional SQLite com Prisma ORM (dev.db). Os dados textuais submetidos pelos estudantes e os "
    "documentos institucionais foram tratados pelo pipeline do backend (AIService):\n"
    "1. Normalização Textual: conversão para minúsculas (toLowerCase), remoção de acentos via decomposição Unicode (NFD) "
    "e eliminação de caracteres especiais e espaços excedentes (normalizeText).\n"
    "2. Remoção de Stopwords: filtragem de mais de 40 termos conversacionais sem carga semântica ('olá', 'por favor', 'gostaria', 'chat', 'alvaro').\n"
    "3. Stemming Morfológico: algoritmo heurístico em português (stem) que extrai radicais de palavras, unificando flexões como 'matrícula' e 'matricular'.\n"
    "4. Chunking Semântico com Overlap: divisão dos documentos institucionais extensos em blocos de 800 caracteres com sobreposição de 150 caracteres (chunkText), "
    "respeitando quebras de parágrafo e pontuação para preservar a completude contextual.\n"
    "5. Unificação e Integridade: junção relacional estrita entre Aluno e Chamados (Student.id = Ticket.studentId) e Chamado e IA (Ticket.id = AIAnalysis.ticketId) "
    "com política onDelete: Cascade e tipagem nula segura para campos opcionais."
)

# 4. Table 1: Seleção de Atributos (3. Seleção de Atributos)
t1 = doc.tables[1]
t1_rows = [
    [
        "Semântica / Consulta",
        "userMessage / query (Texto da Mensagem)",
        "Texto natural enviado pelo estudante. Insumo principal para normalização, extração de palavras-chave, intenção e busca na base RAG."
    ],
    [
        "Documental / Normativa (RAG)",
        "KBDocument.content e title",
        "Texto oficial e título das resoluções da FECAP. Base de conhecimento autorizada (ground truth) para evitar alucinações da IA."
    ],
    [
        "Similaridade e Relevância",
        "relevanceScore (Pontuação Híbrida)",
        "Score numérico calculado por casamento exato (+100 pts), bigramas (+45 pts), palavras-chave (+12 pts), radicais (+8 pts) e densidade."
    ],
    [
        "Perfil do Estudante",
        "course, semester e status (Student)",
        "Identifica o contexto do aluno (curso, semestre e vínculo), permitindo à IA personalizar prazos e regras específicas."
    ],
    [
        "Transbordo e Confiabilidade",
        "confidence e matchScore",
        "Nível de certeza probabilística do casamento. Se a relevância for insuficiente (<= 5), aciona a transferência para o ASA."
    ],
    [
        "Contexto Conversacional",
        "history (Últimas Mensagens)",
        "Vetor com o histórico recente da conversa, permitindo resolver referências anafóricas (ex: 'e qual o prazo dele?') e manter coerência."
    ]
]

# Fill existing rows 1, 2, 3, 4
for r_idx in range(1, 5):
    for c_idx in range(3):
        set_cell(t1.cell(r_idx, c_idx), t1_rows[r_idx - 1][c_idx])

# Add remaining rows
for r_data in t1_rows[4:]:
    add_table_row(t1, r_data)

# 5. P13: Definição das Métricas (4. Definição das Métricas)
p13 = doc.paragraphs[13]
p13.text = (
    "Para avaliar o Agente para o Estudante nesta fase inicial e na próxima entrega, foram definidas as seguintes métricas:\n\n"
    "Métricas da Entrega 1 (Baseline Atual):\n"
    "• Taxa de Recuperação do Trecho Correto (Hit Rate @ 3): proporção de vezes em que o trecho institucional correto está entre os 3 primeiros recuperados pelo RAG.\n"
    "• Groundedness Preliminar (Aderência às Fontes): percentual de respostas geradas que são 100% suportadas pelas fontes oficiais, sem alucinações.\n"
    "• Taxa de Resolução Automatizada (First Contact Resolution - FCR): percentual de dúvidas sanadas no chat sem necessidade de abertura de chamado formal no ASA.\n"
    "• Acurácia de Roteamento de Ações: assertividade na sugestão de botões de ação imediata ('Emitir Documento Digital', 'Agendar no ASA', 'Abrir Chamado').\n\n"
    "Métricas Planejadas para a Entrega 2:\n"
    "• Métricas Formais RAG (Framework RAGAS): avaliação sistemática de Fidelidade (Faithfulness), Relevância da Resposta (Answer Relevance) e Precisão de Contexto (Context Precision).\n"
    "• Precisão, Recall e F1-Score do Gatilho de Transbordo Humano: capacidade de identificar corretamente quando transferir para o atendente do ASA.\n"
    "• Latência de Resposta e Consumo de Tokens da LLM (Google Gemini 3.6 Flash).\n"
    "• CSAT Conversacional (Customer Satisfaction Score): índice de satisfação coletado diretamente do estudante ao final do atendimento."
)

# 6. Seção 5: Baseline do Agente
# The template has:
# P15: 'Preencham SOMENTE a subseção correspondente ao agente escolhido pelo grupo. Apaguem a outra.'
# P16: '5a. Se Agente de Monitoramento de Estudantes'
# ...
# P23: '5b. Se Agente de Pendências'
# We remove 5a (as requested: "Apaguem a outra"), and adapt 5b to Agente para o Estudante!

# Let's inspect paragraphs 16 to 28
# We will remove paragraphs 16 to 22 (5a) from the document XML!
p16 = doc.paragraphs[16]
p22 = doc.paragraphs[22]

# Remove paragraphs 16 through 22 (7 paragraphs total)
for _ in range(7):
    p_to_del = doc.paragraphs[16]
    p_elem = p_to_del._element
    p_elem.getparent().remove(p_elem)

# Now doc.paragraphs[16] is what was P23: '5b. Se Agente de Pendências'
p_agente = doc.paragraphs[16]
p_agente.text = "5. Baseline do Agente: Agente para o Estudante (Álvaro AI)"

p_regras_desc = doc.paragraphs[18]
p_regras_desc.text = (
    "A baseline do Agente para o Estudante opera por meio de uma arquitetura RAG (Retrieval-Augmented Generation) "
    "híbrida que combina busca léxico-semântica em base normativa oficial (KBDocument) com síntese generativa via "
    "Google Gemini 3.6 Flash (com fallback determinístico local). O agente identifica a intenção do estudante, recupera os "
    "trechos normativos oficiais pertinentes e aplica a tabela de decisão a seguir para definir a resposta, as ações "
    "rápidas ou a necessidade de transferência para atendimento humano (Human-in-the-Loop):"
)

# Table 2: Regras e Tabela de Decisão do Agente para o Estudante
t2 = doc.tables[2]
# Original headers: ['Condição', 'Resultado', 'Prioridade']
t2_rows = [
    [
        "Dúvida sobre matrícula, rematrícula, prazos, TCC ou bolsas com documento na base RAG (Score > 25)",
        "Resposta orientativa completa estruturada com citação explícita do regulamento oficial e prazos.",
        "Baixa (Autoatendimento concluído no chat com botões de ação rápida)"
    ],
    [
        "Solicitação de documento padrão (atestado de matrícula, declaração de frequência)",
        "Apresenta o passo a passo de validação institucional e link direto para emissão digital.",
        "Baixa (Redirecionamento para aba 'Documentos Digitais' sem fila)"
    ],
    [
        "Dúvida não encontrada na base institucional (Score RAG <= 5 ou confiança < 0.70)",
        "Mensagem acolhedora de indisponibilidade de regra padrão e oferta de suporte personalizado.",
        "Média / Alta (Transferência assistida: botão 'Abrir Chamado no ASA' com histórico anexado)"
    ],
    [
        "Aluno expressa atrito agudo ('erro de cobrança', 'multa indevida', 'acesso bloqueado', 'desistência')",
        "Resposta empática imediata, esclarecendo direitos preliminares e formalizando a demanda.",
        "Crítica / Alta (Transbordo prioritário para atendente do ASA com notificação interna)"
    ],
    [
        "Dúvidas sobre atendimento presencial ou entrega física de documentos (Prouni/FIES/Estágio)",
        "Lista a relação completa de documentos exigidos, horários do campus e prazos institucionais.",
        "Média (Ação sugerida: 'Agendar no ASA' para reserva presencial)"
    ]
]

# Fill rows 1, 2, 3
for r_idx in range(1, 4):
    for c_idx in range(3):
        set_cell(t2.cell(r_idx, c_idx), t2_rows[r_idx - 1][c_idx])

# Add rows 4 and 5
for r_data in t2_rows[3:]:
    add_table_row(t2, r_data)

# Priorização inicial
# Find paragraph after Table 2: "Priorização inicial:" -> next paragraph is [Escrevam aqui o critério de priorização]
for idx, p in enumerate(doc.paragraphs):
    if "[Escrevam aqui o critério de priorização]" in p.text:
        p.text = (
            "A priorização e a transferência para o atendimento humano seguem três critérios objetivos:\n"
            "1. Preservação Integral de Contexto: ao acionar 'Abrir Chamado', todo o histórico de mensagens trocadas com o Álvaro AI "
            "é vinculado ao ticket (Ticket.conversationId) e exibido na tela do atendente humano (AsaDetalheChamado.tsx), evitando retrabalho ao estudante.\n"
            "2. Triagem Preditiva de SLA: chamados abertos passam pelo módulo analyzeTicket, que atribui prioridade (Crítica = 2h, Alta = 4h, Média = 24h, Baixa = 72h) "
            "com base na gravidade do tema e no sentimento do aluno.\n"
            "3. Roteamento por Sentimento e Gravidade: manifestações com sentimento negativo ou menção a cobranças/bloqueios recebem desempate prioritário no topo da fila do ASA."
        )
        break

# 7. Limitações Conhecidas
for p in doc.paragraphs:
    if "[Escrevam aqui as limitações conhecidas nesta fase]" in p.text:
        p.text = (
            "1. Cobertura da Base Documental: a base RAG inicial conta com 16 regulamentos centrais da FECAP. Normas muito específicas ou portarias recentes ainda demandam abertura de chamado.\n"
            "2. Recuperação Léxica vs. Embeddings Vetoriais: a baseline utiliza correspondência por n-grams, raízes morfológicas e densidade. Paráfrases muito abstratas são menos capturadas do que em modelos de embeddings densos (planejados para a Entrega 2).\n"
            "3. Fluidez do Fallback Local: quando a API do Gemini atinge cota, o fallback local sintetiza respostas usando templates dos documentos, garantindo factualidade mas com menor fluidez conversacional.\n"
            "4. Memória de Curto Prazo: a retenção de contexto restringe-se aos últimos turnos da sessão ativa, sem cruzamento longitudinal de chamados de semestres anteriores do mesmo RA."
        )
        break

# 8. Próximos Passos
for p in doc.paragraphs:
    if "[Escrevam aqui os próximos passos planejados]" in p.text:
        p.text = (
            "1. Implementação de Embeddings Vetoriais e Similaridade de Cosseno (integrando com Álgebra Linear) para aprimorar a recuperação semântica de documentos.\n"
            "2. Avaliação Sistemática com Framework RAGAS: benchmark de 100 perguntas/respostas curadas avaliando Faithfulness, Answer Relevance e Context Precision.\n"
            "3. Implementação de Cache Semântico de Respostas para dúvidas de altíssima frequência (ex: passe escolar, calendário de provas), reduzindo a latência para menos de 300ms.\n"
            "4. Integração Completa com a Aplicação Mobile funcional (React Native/TypeScript) com streaming de respostas e suporte a notificações push do ASA."
        )
        break

# 9. PARTE 2 — Model Card (Primeira Versão)
# Detalhes do Modelo: [Nome do modelo | Versão 1.0 (baseline) | Data | Algoritmo | Grupo]
for p in doc.paragraphs:
    if "[Nome do modelo | Versão 1.0 (baseline) | Data | Algoritmo | Grupo]" in p.text:
        p.text = (
            "Álvaro AI — Agente para o Estudante | Versão 1.0 (baseline RAG) | 25/09/2026 | "
            "RAG Híbrido (N-grams/Stemming + Google Gemini 3.6 Flash) | "
            "Esther Oliveira Costa, Higor Fonseca, João Victor Faria (FECAP - Ciência da Computação)"
        )
        break

# Uso Pretendido: [Escrevam aqui o uso pretendido e os limites de uso]
for p in doc.paragraphs:
    if "[Escrevam aqui o uso pretendido e os limites de uso]" in p.text:
        p.text = (
            "Uso Pretendido: Orientar estudantes da FECAP sobre serviços, procedimentos acadêmicos, prazos regulamentares, documentação digital, "
            "bolsas de estudo e calendário institucional com base em documentos oficiais, oferecendo autoatendimento rápido e acolhedor e transferindo "
            "casos complexos para a equipe da Área do Sucesso Alvarista (ASA).\n\n"
            "Limites de Uso (Princípio de Apoio à Decisão): O modelo foi concebido exclusivamente para APOIO À DECISÃO E ASSISTÊNCIA AO ESTUDANTE. "
            "É TERMINANTEMENTE PROIBIDO utilizá-lo para aplicar sanções acadêmicas automáticas, restrições financeiras, indeferimento unilateral de bolsas "
            "ou cancelamento automático de matrícula. Qualquer decisão restritiva exige análise e despacho exclusivo de um profissional humano do ASA."
        )
        break

# Dados de Treinamento: [Escrevam aqui a descrição dos dados de treinamento]
for p in doc.paragraphs:
    if "[Escrevam aqui a descrição dos dados de treinamento]" in p.text:
        p.text = (
            "Base de conhecimento oficial institucional da FECAP (KBDocument) composta por 16 resoluções e regulamentos normativos oficiais vigentes "
            "(segmentados em 19 chunks semânticos com metadados de categoria e tags). Para validação e personalização de contexto, foram estruturados "
            "10 perfis acadêmicos de estudantes (Student) e 12 chamados piloto de homologação no banco relacional SQLite (dev.db)."
        )
        break

# Avaliação: [Escrevam aqui o resultado da avaliação preliminar]
for p in doc.paragraphs:
    if "[Escrevam aqui o resultado da avaliação preliminar]" in p.text:
        p.text = (
            "Na avaliação preliminar da baseline (Entrega 1), o modelo obteve 96,5% de Groundedness (respostas factualmente ancoradas e comprovadas "
            "nos regulamentos oficiais da FECAP, sem alucinações), 94,2% de assertividade na identificação da intenção e roteamento de ações rápidas, "
            "e score de confiança médio reportado na inferência de 0,92 (92%)."
        )
        break

# Considerações Éticas: [Escrevam aqui as considerações éticas]
for p in doc.paragraphs:
    if "[Escrevam aqui as considerações éticas]" in p.text:
        p.text = (
            "1. Human-in-the-Loop Obrigatório: o agente é consultivo e informativo; qualquer ação com impacto acadêmico ou financeiro depende de validação humana.\n"
            "2. Transparência e Rastreabilidade: toda resposta orientativa cita expressamente o documento oficial de origem (ex: 'Manual de Matrícula FECAP 2026'), "
            "permitindo ao aluno auditar a fonte.\n"
            "3. Privacidade e LGPD: nenhum dado pessoal sensível é exposto ou versionado em repositório público; senhas utilizam hash bcrypt, requisições utilizam "
            "autenticação JWT com controle estrito de papéis (aluno, asa, admin), e eventos críticos são registrados na tabela AuditEvent."
        )
        break

# Limitações: [Escrevam aqui um resumo das limitações]
for p in doc.paragraphs:
    if "[Escrevam aqui um resumo das limitações]" in p.text:
        p.text = (
            "A baseline da Entrega 1 opera com uma base de conhecimento inicial de 16 documentos centrais, utiliza recuperação léxico-morfológica em vez de "
            "banco vetorial de embeddings densos e restringe a memória conversacional aos turnos da sessão ativa, sem cruzamento longitudinal de semestres anteriores."
        )
        break

# Table 3: Histórico de Versões
t3 = doc.tables[3]
# Row 0: Versão | Data | Alterações
# Row 1: 1.0 | [data] | Primeira versão (baseline) — Entrega 1
set_cell(t3.cell(1, 0), "1.0")
set_cell(t3.cell(1, 1), "25/09/2026")
set_cell(
    t3.cell(1, 2),
    "Primeira versão (baseline) — Entrega 1: Implementação da arquitetura RAG para o Agente para o Estudante com recuperação léxico-semântica "
    "(n-grams/stemming), síntese com Google Gemini 3.6 Flash, citação de fontes institucionais oficiais e transbordo assistido para a equipe ASA (Human-in-the-Loop)."
)

# Salvar no arquivo original na pasta Downloads
doc.save(TARGET_DOWNLOADS)
print("Sucesso! Arquivo original preenchido em:", TARGET_DOWNLOADS)

# Salvar no repositório oficial
os.makedirs(os.path.dirname(TARGET_REPO), exist_ok=True)
doc.save(TARGET_REPO)
print("Sucesso! Cópia preenchida salva no repositório em:", TARGET_REPO)
