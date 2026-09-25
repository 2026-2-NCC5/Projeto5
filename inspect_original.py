import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'C:\Users\24026811\Downloads\PI_Template_Entrega1_Original_Backup.docx')

print("=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    print(f"P{i:02d} [style={p.style.name}]: {repr(p.text)}")

print("\n=== TABLES ===")
for t_idx, t in enumerate(doc.tables):
    print(f"\n--- Table {t_idx} ({len(t.rows)} rows x {len(t.columns)} cols) ---")
    for r_idx, r in enumerate(t.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
        print(f"  R{r_idx}: {cells}")
